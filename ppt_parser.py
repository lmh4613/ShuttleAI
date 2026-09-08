import os
import json
import io
import re
import streamlit as st
from dotenv import load_dotenv
from google import genai
from google.genai import types
from pptx import Presentation
from pdf2image import convert_from_bytes

load_dotenv()

def get_gemini_client():
    api_key = os.getenv("GEMINI_API_KEY")
    if not api_key:
        try:
            api_key = st.secrets.get("GEMINI_API_KEY")
        except Exception:
            api_key = None
            
    if not api_key:
        raise ValueError(".env 파일이나 Streamlit Secrets에 GEMINI_API_KEY가 설정되어 있지 않습니다.")
    return genai.Client(api_key=api_key)

def parse_shuttle_document(uploaded_file):
    client = get_gemini_client()
    file_name = uploaded_file.name.lower()
    
    prompt = """
    너는 셔틀버스 노선도 분석 전문가야. 제공된 파일/이미지 내용을 분석해서 노선 및 정류장 정보를 정확하게 추출해 줘.
    
    [핵심 규칙]
    1. 동일한 시간대나 목적지라도 (1호), (2호), (3호) 또는 A, B, C 등 세부 노선이나 차량 번호가 나뉘어져 있다면, 이를 절대 합치지 말고 각각 독립된 노선으로 분리해서 추출할 것.
    2. `route_name`에는 반드시 세부 노선 번호나 이름까지 포함할 것 (예: "(출근)망포/영통 1호", "(퇴근)병점/서천").
    3. 각 노선 및 정류장의 위치가 서울 지역(예: 강남, 사당 등)인지 경기 지역(예: 수원, 판교, 동탄, 용인 등)인지 판단하여 `region` 필드에 정확히 "seoul" 또는 "gyeonggi"를 입력할 것.
    4. 반드시 아래 JSON 배열 형식으로만 응답해야 해. 마크다운(` ```json `) 표기 없이 pure JSON 문자열만 출력해.
    
    [응답 데이터 형식]
    [
      {
        "route_name": "(출근)망포/영통 1호",
        "stop_name": "정자역 3번출구",
        "arrival_time": "08:15",
        "region": "gyeonggi",
        "address": "위치 설명 또는 주소"
      }
    ]
    """

    contents = []

    if file_name.endswith('.pdf'):
        images = convert_from_bytes(uploaded_file.read())
        for img in images:
            img_byte_arr = io.BytesIO()
            img.save(img_byte_arr, format='JPEG', quality=85)
            contents.append(
                types.Part.from_bytes(
                    data=img_byte_arr.getvalue(),
                    mime_type="image/jpeg"
                )
            )
        contents.append(prompt)

    elif file_name.endswith('.pptx'):
        prs = Presentation(uploaded_file)
        text_content = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    text_content.append(shape.text_frame.text)
                if shape.has_table:
                    for row in shape.table.rows:
                        text_content.append(" | ".join([cell.text.strip() for cell in row.cells]))
        
        extracted_text = "\n".join(text_content)
        contents = [f"PPTX extracted content:\n{extracted_text}", prompt]
    
    else:
        raise ValueError("지원하지 않는 파일 형식입니다. (PDF, PPTX만 가능)")

    response = client.models.generate_content(
        model='gemini-1.5-flash',
        contents=contents
    )
    
    raw_text = response.text.strip()
    if "```" in raw_text:
        raw_text = re.sub(r'```(?:json)?\s*', '', raw_text)
        raw_text = raw_text.replace('```', '').strip()
    
    return json.loads(raw_text)